#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import math
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt
except ImportError as exc:  # pragma: no cover
    raise SystemExit("python-docx não instalado. Execute o pipeline pelo ambiente virtual do Canopy2Ground.") from exc

try:
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
except ImportError as exc:  # pragma: no cover
    raise SystemExit("reportlab não instalado. Execute o pipeline pelo ambiente virtual do Canopy2Ground.") from exc

ROOT_TITLE = "Canopy2Ground - Relatório automático de resultados"
MONTHS_PT = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio", 6: "junho",
    7: "julho", 8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro",
}


def read_csv(path: Path) -> pd.DataFrame | None:
    return pd.read_csv(path, low_memory=False) if path.exists() else None


def read_json(path: Path) -> dict | None:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def fmt(x, digits=3, decimal_comma=True):
    if x is None:
        return "NA"
    try:
        x = float(x)
    except Exception:
        return str(x)
    if not math.isfinite(x):
        return "NA"
    s = f"{x:.{digits}f}"
    return s.replace(".", ",") if decimal_comma else s


def fmt_signed(x, digits=3):
    try:
        x = float(x)
    except Exception:
        return str(x)
    if not math.isfinite(x):
        return "NA"
    s = f"{x:+.{digits}f}"
    return s.replace(".", ",")


def fmt_int(x) -> str:
    try:
        return f"{int(x):,}".replace(",", ".")
    except Exception:
        return str(x)


def fmt_p(p: float) -> str:
    p = float(p)
    return "p < 0,001" if p < 0.001 else f"p = {fmt(p, 3)}"


def significance_label(p: float, alpha: float = 0.05) -> str:
    return "significativa" if float(p) < alpha else "não significativa"


def significance_clause(p: float, alpha: float = 0.05) -> str:
    return "foi significativa" if float(p) < alpha else "não foi significativa"


def stat_row(stats: pd.DataFrame, variable: str, contrast: str) -> pd.Series:
    z = stats[(stats["variable"] == variable) & (stats["contrast"] == contrast)]
    if z.empty:
        raise KeyError(f"Resultado ausente: {variable} / {contrast}")
    return z.iloc[0]


def model_var(model: dict, variable: str) -> dict:
    return model.get("variables", {}).get(variable, {})


def reduction(cv: dict) -> float:
    try:
        raw = float(cv.get("raw_RMSE")); lin = float(cv.get("linear_RMSE"))
        return 100.0 * (raw - lin) / raw if raw else float("nan")
    except Exception:
        return float("nan")


def _difference_sentence(source: str, row: pd.Series, unit: str) -> str:
    p = float(row["p_holm_within_variable"])
    return (
        f"A diferença média {source} − pastagem foi de {fmt_signed(row['bias_ar1'], 3)} {unit} e "
        f"{significance_clause(p)} segundo o modelo GLSAR AR(1), após correção de Holm ({fmt_p(p)})."
    )


def build_sections(root: Path) -> list[tuple[str, list[str], int | None]]:
    micro = root / "results" / "validation" / "climate" / "microclimate"
    climate = root / "results" / "validation" / "climate"
    harmon = root / "results" / "harmonization"

    stats = read_csv(micro / "microclimate_source_comparison_all.csv")
    tri_daily = read_csv(micro / "microclimate_common_daily_all.csv")
    tri_hourly = read_csv(micro / "microclimate_common_hourly_all.csv")
    pair_daily = read_csv(micro / "era5_field_common_daily_all.csv")
    model = read_json(harmon / "current_model.json")
    precip = read_csv(climate / "best_temporal_alignment.csv")
    pmonth = read_csv(climate / "precip_monthly_paired.csv")

    sections: list[tuple[str, list[str], int | None]] = []

    # 3.1 — always written from the current three-source comparison table.
    p31: list[str] = []
    if stats is not None and tri_daily is not None:
        n_days = len(tri_daily)
        n_hours = len(tri_hourly) if tri_hourly is not None else int(pd.to_numeric(tri_daily.get("n_common_hours"), errors="coerce").sum())
        p31.append(
            f"Foram obtidos {fmt_int(n_hours)} registros horários simultâneos entre a pastagem, o ERA5-Land e o INMET, "
            f"correspondentes a {fmt_int(n_days)} dias com cobertura diária. As três fontes apresentaram comportamento temporal semelhante "
            f"para temperatura, umidade relativa e déficit de pressão de vapor (VPD), embora com diferenças na magnitude dos valores observados."
        )

        t_e = stat_row(stats, "temperature", "ERA5 - FIELD")
        t_i = stat_row(stats, "temperature", "INMET - FIELD")
        era5_temp = _difference_sentence("ERA5-Land", t_e, "°C")
        era5_temp = era5_temp[0].lower() + era5_temp[1:]
        p31.append(
            "Para temperatura do ar, " + era5_temp + " "
            + _difference_sentence("INMET", t_i, "°C") + " "
            + f"A correlação com a pastagem foi de {fmt(t_e['r'], 3)} para ERA5-Land e {fmt(t_i['r'], 3)} para INMET."
        )

        rh_e = stat_row(stats, "rh", "ERA5 - FIELD")
        rh_i = stat_row(stats, "rh", "INMET - FIELD")
        p31.append(
            f"A umidade relativa média foi de {fmt(rh_e['mean_B'], 1)}% na pastagem, {fmt(rh_e['mean_A'], 1)}% no ERA5-Land e "
            f"{fmt(rh_i['mean_A'], 1)}% no INMET. As diferenças médias foram de {fmt_signed(rh_e['bias_ar1'], 2)} pontos percentuais "
            f"para ERA5-Land ({significance_label(rh_e['p_holm_within_variable'])}; {fmt_p(rh_e['p_holm_within_variable'])}) e "
            f"{fmt_signed(rh_i['bias_ar1'], 2)} pontos percentuais para INMET ({significance_label(rh_i['p_holm_within_variable'])}; "
            f"{fmt_p(rh_i['p_holm_within_variable'])}). As correlações foram de {fmt(rh_e['r'], 3)} e {fmt(rh_i['r'], 3)}, respectivamente, "
            f"enquanto o CCC entre ERA5-Land e pastagem foi de {fmt(rh_e['ccc'], 3)}."
        )

        v_e = stat_row(stats, "vpd", "ERA5 - FIELD")
        v_i = stat_row(stats, "vpd", "INMET - FIELD")
        p31.append(
            f"O VPD médio registrado na pastagem foi de {fmt(v_e['mean_B'], 3)} kPa, comparado a {fmt(v_e['mean_A'], 3)} kPa no ERA5-Land "
            f"e {fmt(v_i['mean_A'], 3)} kPa no INMET. As diferenças em relação à pastagem foram de {fmt_signed(v_e['bias_ar1'], 3)} kPa "
            f"para ERA5-Land ({significance_label(v_e['p_holm_within_variable'])}; {fmt_p(v_e['p_holm_within_variable'])}) e "
            f"{fmt_signed(v_i['bias_ar1'], 3)} kPa para INMET ({significance_label(v_i['p_holm_within_variable'])}; "
            f"{fmt_p(v_i['p_holm_within_variable'])}). As correlações com a pastagem foram de {fmt(v_e['r'], 3)} e {fmt(v_i['r'], 3)}, "
            f"respectivamente, enquanto o CCC entre ERA5-Land e pastagem foi de {fmt(v_e['ccc'], 3)}. A correlação elevada, portanto, "
            f"não implicou concordância equivalente entre as magnitudes das séries (Figura 1)."
        )
    else:
        p31.append("A comparação simultânea Pastagem + ERA5-Land + INMET não estava disponível na execução atual.")
    sections.append(("3.1 Representação das condições atmosféricas locais", p31, 1))

    # 3.2 — dynamic ERA5 harmonization.
    p32: list[str] = []
    if model is not None and pair_daily is not None:
        n_pair = int(model.get("validation_design", {}).get("n_days", len(pair_daily)))
        period_start = model.get("validation_design", {}).get("period_start", "")
        period_end = model.get("validation_design", {}).get("period_end", "")
        p32.append(
            f"As regressões entre ERA5-Land e as observações na pastagem foram atualizadas com {fmt_int(n_pair)} dias pareados "
            f"({period_start} a {period_end}). Temperatura, umidade relativa e VPD apresentaram relações lineares consistentes, "
            f"com ganhos distintos após a harmonização (Figura 2)."
        )
        temp = model_var(model, "temperature"); rh = model_var(model, "rh"); vpd = model_var(model, "vpd")
        p32.append(
            f"Os modelos lineares correntes foram T_pasto = {fmt(temp.get('intercept'), 3)} + {fmt(temp.get('slope'), 4)}(T_ERA5), "
            f"RH_pasto = {fmt(rh.get('intercept'), 3)} + {fmt(rh.get('slope'), 4)}(RH_ERA5) e "
            f"VPD_pasto = {fmt(vpd.get('intercept'), 4)} + {fmt(vpd.get('slope'), 4)}(VPD_ERA5). "
            f"Os coeficientes são reestimados cumulativamente quando novos dias pareados são incorporados."
        )
        tcv = temp.get("cross_validation", {}); rcv = rh.get("cross_validation", {}); vcv = vpd.get("cross_validation", {})
        p32.append(
            f"Na validação cruzada leave-one-day-out, com embargo temporal de um dia, o RMSE da temperatura diminuiu de "
            f"{fmt(tcv.get('raw_RMSE'), 3)} para {fmt(tcv.get('linear_RMSE'), 3)} °C ({fmt(reduction(tcv), 1)}% de redução), "
            f"e o CCC passou de {fmt(tcv.get('raw_CCC'), 3)} para {fmt(tcv.get('linear_CCC'), 3)}. Para umidade relativa, "
            f"o RMSE diminuiu de {fmt(rcv.get('raw_RMSE'), 2)} para {fmt(rcv.get('linear_RMSE'), 2)} pontos percentuais "
            f"({fmt(reduction(rcv), 1)}%), com CCC de {fmt(rcv.get('raw_CCC'), 3)} para {fmt(rcv.get('linear_CCC'), 3)}. "
            f"Para VPD, o RMSE diminuiu de {fmt(vcv.get('raw_RMSE'), 3)} para {fmt(vcv.get('linear_RMSE'), 3)} kPa "
            f"({fmt(reduction(vcv), 1)}%), e o CCC passou de {fmt(vcv.get('raw_CCC'), 3)} para {fmt(vcv.get('linear_CCC'), 3)}."
        )
        p32.append(
            "Com base no desempenho fora da amostra, a versão operacional aplica regressão linear às três variáveis atmosféricas. "
            "Os valores brutos do ERA5-Land permanecem preservados ao lado das séries harmonizadas (Figura 3)."
        )
    else:
        p32.append("A harmonização não foi executada nesta atualização por ausência de pareamento diário Pastagem + ERA5-Land.")
    sections.append(("3.2 Harmonização do ERA5-Land", p32, 2))
    sections.append(("", [], 3))

    # 3.3 — precipitation: temporal + monthly comparison only, no regression.
    p33: list[str] = []
    if precip is not None:
        pp = precip[precip["source"].isin(["ERA5-Land", "INMET"])].copy()
        by = {r.source: r for r in pp.itertuples(index=False)}
        if "ERA5-Land" in by and "INMET" in by:
            e, i = by["ERA5-Land"], by["INMET"]
            p33.append(
                f"A representação da precipitação diária foi mais variável do que a observada para as condições atmosféricas. "
                f"O ERA5-Land apresentou RMSE de {fmt(e.RMSE, 2)} mm dia⁻¹, correlação de {fmt(e.r, 3)} e PBIAS de {fmt(e.PBIAS, 1)}%. "
                f"Para o INMET, a correlação foi de {fmt(i.r, 3)}, o RMSE foi de {fmt(i.RMSE, 2)} mm dia⁻¹ e o PBIAS de {fmt(i.PBIAS, 1)}%."
            )
            p33.append(
                "O viés acumulado foi interpretado em conjunto com a distribuição temporal dos erros, pois períodos de superestimativa e "
                "subestimativa podem se compensar ao longo da série."
            )
        elif not pp.empty:
            r = next(pp.itertuples(index=False))
            p33.append(
                f"A validação de precipitação disponível nesta execução incluiu {r.source}, com RMSE de {fmt(r.RMSE, 2)} mm dia⁻¹, "
                f"correlação de {fmt(r.r, 3)} e PBIAS de {fmt(r.PBIAS, 1)}%."
            )

        if pmonth is not None and not pp.empty:
            chunks = []
            for row in pp.itertuples(index=False):
                z = pmonth[(pmonth["source"] == row.source) & (pmonth["window"] == row.window) & (pmonth["shift_days"] == row.shift_days)].copy()
                if not z.empty:
                    chunks.append(z)
            if chunks:
                m = pd.concat(chunks, ignore_index=True)
                m["field_coverage"] = pd.to_numeric(m["n_field_valid_days"], errors="coerce") / pd.to_numeric(m["calendar_days"], errors="coerce")
                m["paired_field_coverage"] = pd.to_numeric(m["paired_field_coverage"], errors="coerce")
                mc = m[(m["field_coverage"] >= 0.90) & (m["paired_field_coverage"] >= 0.90)].copy()
                if not mc.empty:
                    mc["abs_error"] = (pd.to_numeric(mc["product_rain_mm"], errors="coerce") - pd.to_numeric(mc["field_rain_mm"], errors="coerce")).abs()
                    rank = mc.groupby("month", as_index=False)["abs_error"].mean().sort_values("abs_error", ascending=False)
                    months = rank["month"].head(2).tolist()
                else:
                    months = []
                examples = []
                for month in months:
                    z = mc[mc["month"] == month]
                    dt = pd.to_datetime(month + "-01")
                    field_val = float(z["field_rain_mm"].iloc[0])
                    vals = {r.source: float(r.product_rain_mm) for r in z.itertuples(index=False)}
                    parts = [f"{MONTHS_PT[dt.month]}: {fmt(field_val,1)} mm na pastagem"]
                    if "ERA5-Land" in vals:
                        parts.append(f"{fmt(vals['ERA5-Land'],1)} mm no ERA5-Land")
                    if "INMET" in vals:
                        parts.append(f"{fmt(vals['INMET'],1)} mm no INMET")
                    examples.append(", ".join(parts))
                if examples:
                    p33.append(
                        "Na escala mensal, as diferenças também foram evidentes; entre os meses com pelo menos 90% de cobertura, destacaram-se "
                        + "; e ".join(examples) + " (Figura 4)."
                    )
    if not p33:
        p33.append("A comparação de precipitação com ERA5-Land e INMET não estava disponível na execução atual.")
    sections.append(("3.3 Representação da precipitação diária", p33, 4))
    return sections


def figure_caption(n: int) -> str:
    caps = {
        1: "Figura 1. Distribuição das médias diárias de (A) temperatura do ar, (B) umidade relativa e (C) déficit de pressão de vapor (VPD) observadas na pastagem e obtidas pelo ERA5-Land e INMET. Letras diferentes indicam diferenças entre as fontes com base em contrastes diários pareados ajustados por GLSAR AR(1), com correção de Holm para comparações múltiplas.",
        2: "Figura 2. Relação entre valores diários de (A) temperatura do ar, (B) umidade relativa e (C) déficit de pressão de vapor (VPD) obtidos pelo ERA5-Land e observados na pastagem. A linha contínua representa a regressão linear ajustada e a linha tracejada indica a relação 1:1. São apresentadas as equações, os coeficientes de determinação (R²) e o número de dias pareados.",
        3: "Figura 3. Desempenho do ERA5-Land bruto e da regressão linear na validação leave-one-day-out com embargo temporal de um dia. São apresentados (A) a redução percentual do RMSE e (B) o coeficiente de concordância de Lin (CCC).",
        4: "Figura 4. Comparação da precipitação observada na pastagem com ERA5-Land e INMET. (A) precipitação diária para o alinhamento temporal selecionado; (B) totais mensais para meses com pelo menos 90% de cobertura dos dias válidos no campo e 90% de cobertura pareada do respectivo produto.",
    }
    return caps[n]


def figure_path(root: Path, n: int) -> Path:
    names = {
        1: "figure_01_microclimate_sources.png",
        2: "figure_02_era5_pasture_regressions.png",
        3: "figure_03_harmonization_performance.png",
        4: "figure_04_precipitation_comparison.png",
    }
    return root / "results" / "figures" / names[n]


def add_docx_figure(doc: Document, root: Path, n: int) -> None:
    p = figure_path(root, n)
    if not p.exists():
        return
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(p), width=Cm(16.5))
    cap = doc.add_paragraph(figure_caption(n)); cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if cap.runs:
        cap.runs[0].font.size = Pt(9)


def write_docx(root: Path, sections, out: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0); sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"; styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Times New Roman"; styles["Title"].font.size = Pt(14); styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Times New Roman"; styles["Heading 1"].font.size = Pt(11); styles["Heading 1"].font.bold = True
    t = doc.add_paragraph(style="Title"); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; t.add_run(ROOT_TITLE)
    meta = doc.add_paragraph(f"Gerado automaticamente em {datetime.now().strftime('%d/%m/%Y %H:%M')}. Os valores refletem os produtos da execução corrente.")
    meta.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for title, paras, fig in sections:
        if title:
            doc.add_heading(title, level=1)
        for txt in paras:
            p = doc.add_paragraph(txt); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if fig is not None:
            add_docx_figure(doc, root, fig)
    doc.save(out)


def write_pdf(root: Path, sections, out: Path) -> None:
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleC2G", parent=styles["Title"], fontName="Times-Bold", fontSize=14, leading=17, alignment=1, spaceAfter=12)
    head = ParagraphStyle("HeadC2G", parent=styles["Heading2"], fontName="Times-Bold", fontSize=11, leading=13, spaceBefore=8, spaceAfter=5)
    body = ParagraphStyle("BodyC2G", parent=styles["BodyText"], fontName="Times-Roman", fontSize=9.5, leading=12, alignment=TA_JUSTIFY, spaceAfter=6)
    cap = ParagraphStyle("CapC2G", parent=body, fontSize=8.5, leading=10, spaceAfter=10)
    doc = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=1.8*cm, leftMargin=1.8*cm, topMargin=1.8*cm, bottomMargin=1.8*cm)
    story = [Paragraph(ROOT_TITLE, title), Paragraph(f"Gerado automaticamente em {datetime.now().strftime('%d/%m/%Y %H:%M')}. Os valores refletem os produtos da execução corrente.", body), Spacer(1, 4)]
    max_w = A4[0] - 3.6*cm
    for title_txt, paras, fig in sections:
        if title_txt:
            story.append(Paragraph(title_txt, head))
        for txt in paras:
            safe = txt.replace("&", "&amp;").replace("⁻¹", "<super>-1</super>")
            story.append(Paragraph(safe, body))
        if fig is not None:
            fp = figure_path(root, fig)
            if fp.exists():
                im = Image(str(fp)); ratio = min(max_w / im.imageWidth, (10.5*cm) / im.imageHeight)
                im.drawWidth = im.imageWidth * ratio; im.drawHeight = im.imageHeight * ratio
                story.extend([Spacer(1, 5), im, Paragraph(figure_caption(fig), cap)])
    doc.build(story)


def write_report_inputs(root: Path, out: Path) -> None:
    paths = {
        "three_source_stats": root / "results/validation/climate/microclimate/microclimate_source_comparison_all.csv",
        "three_source_daily": root / "results/validation/climate/microclimate/microclimate_common_daily_all.csv",
        "era5_pasture_daily": root / "results/validation/climate/microclimate/era5_field_common_daily_all.csv",
        "harmonization_model": root / "results/harmonization/current_model.json",
        "precip_alignment": root / "results/validation/climate/best_temporal_alignment.csv",
        "precip_monthly": root / "results/validation/climate/precip_monthly_paired.csv",
    }
    rows = []
    for name, path in paths.items():
        rec = {"input": name, "path": str(path), "exists": path.exists(), "modified_utc": None, "rows": None}
        if path.exists():
            rec["modified_utc"] = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat()
            if path.suffix.lower() == ".csv":
                try:
                    rec["rows"] = len(pd.read_csv(path, low_memory=False))
                except Exception:
                    pass
            elif path.name == "current_model.json":
                try:
                    m = read_json(path) or {}
                    rec["rows"] = m.get("validation_design", {}).get("n_days")
                    rec["model_id"] = m.get("model_id")
                except Exception:
                    pass
        rows.append(rec)
    pd.DataFrame(rows).to_csv(out, index=False, encoding="utf-8-sig")


def main() -> None:
    ap = argparse.ArgumentParser(description="Gera relatório automático a partir dos resultados correntes do Canopy2Ground.")
    ap.add_argument("--root", default=".")
    args = ap.parse_args()
    root = Path(args.root).resolve()
    outdir = root / "results" / "report"; outdir.mkdir(parents=True, exist_ok=True)
    sections = build_sections(root)
    docx_path = outdir / "Canopy2Ground_resultados.docx"
    pdf_path = outdir / "Canopy2Ground_resultados.pdf"
    write_docx(root, sections, docx_path)
    write_pdf(root, sections, pdf_path)
    write_report_inputs(root, outdir / "REPORT_INPUTS.csv")
    print(f"Relatório Word: {docx_path}")
    print(f"Relatório PDF:  {pdf_path}")
    print(f"Entradas auditadas: {outdir / 'REPORT_INPUTS.csv'}")


if __name__ == "__main__":
    main()
