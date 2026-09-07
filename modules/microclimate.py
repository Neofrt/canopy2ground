#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

SITE_DIRS = {"PASTO": "pasto", "SECUNDARIA": "cap", "PRIMARIA": "ref"}


def load_config(path: Path | None):
    return json.loads(path.read_text(encoding="utf-8")) if path and path.exists() else {}


def find_column(columns, pattern):
    rx = re.compile(pattern, re.I)
    for i, c in enumerate(columns):
        if rx.search(str(c)):
            return i
    return None


def workbook_meta(path: Path):
    """Fast metadata read. Also lets us choose only the newest cumulative deployment file."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return {"serial": None, "deployment": None, "has_data": False, "error": str(e)}
    meta = {"serial": None, "deployment": None, "has_data": "Data" in wb.sheetnames, "error": None}
    if "Details" in wb.sheetnames:
        ws = wb["Details"]
        for row in ws.iter_rows(values_only=True):
            if len(row) >= 4:
                k, v = row[2], row[3]
                if k == "Serial Number": meta["serial"] = None if v is None else str(v)
                elif k == "Deployment Number": meta["deployment"] = None if v is None else str(v)
    wb.close()
    return meta


def select_cumulative_files(root: Path):
    candidates=[]
    for folder, site in SITE_DIRS.items():
        d=root/"data"/"raw"/folder/"Termohigrometro"
        if not d.exists(): continue
        for p in d.rglob("*.xlsx"):
            if p.name.startswith("~$"): continue
            meta=workbook_meta(p)
            candidates.append({"path":p,"site":site,"size":p.stat().st_size,**meta})
    # HOBO downloads are cumulative. Keep largest file per site+serial+deployment.
    selected=[]; audits=[]
    groups={}
    for z in candidates:
        if not z["has_data"]:
            audits.append({"site":z["site"],"file":str(z["path"]),"status":"ignored_no_Data_sheet","serial":z["serial"],"deployment":z["deployment"],"note":z["error"] or ""})
            continue
        key=(z["site"],z["serial"] or "unknown",z["deployment"] or "unknown")
        groups.setdefault(key,[]).append(z)
    for key,arr in groups.items():
        arr=sorted(arr,key=lambda z:(z["size"],str(z["path"])),reverse=True)
        selected.append(arr[0])
        for z in arr[1:]:
            audits.append({"site":z["site"],"file":str(z["path"]),"status":"skipped_older_cumulative_download","serial":z["serial"],"deployment":z["deployment"],"note":f"selected={arr[0]['path']}"})
    return selected,audits


def read_one(info):
    path=info["path"]; site=info["site"]
    try:
        wb=load_workbook(path,read_only=True,data_only=True)
        ws=wb["Data"]
        it=ws.iter_rows(values_only=True)
        headers=list(next(it))
        c_dt=find_column(headers,r"Date-Time")
        c_t=find_column(headers,r"^Temperature")
        c_rh=find_column(headers,r"^RH")
        c_par=find_column(headers,r"Photosynthetically Active Radiation")
        c_vpd=find_column(headers,r"Vapor Pressure Deficit")
        c_dew=find_column(headers,r"Dew Point")
        if any(i is None for i in [c_dt,c_t,c_rh,c_par,c_vpd]):
            wb.close(); return None,{"site":site,"file":str(path),"status":"ignored_missing_columns","serial":info["serial"],"deployment":info["deployment"],"note":""}
        rows=[]
        for r in it:
            dt=r[c_dt]
            if dt is None: continue
            rows.append((dt,r[c_t],r[c_rh],r[c_par],r[c_vpd],r[c_dew] if c_dew is not None else None))
        wb.close()
    except Exception as e:
        return None,{"site":site,"file":str(path),"status":"unreadable_Data","serial":info["serial"],"deployment":info["deployment"],"note":str(e)}
    if not rows:
        return None,{"site":site,"file":str(path),"status":"empty_Data","serial":info["serial"],"deployment":info["deployment"],"note":""}
    out=pd.DataFrame(rows,columns=["datetime","temp_C","rh_pct","par_umol","vpd_kPa","dewpoint_C"])
    out["datetime"]=pd.to_datetime(out["datetime"],errors="coerce")
    for c in ["temp_C","rh_pct","par_umol","vpd_kPa","dewpoint_C"]: out[c]=pd.to_numeric(out[c],errors="coerce")
    out=out[out["datetime"].notna()].copy(); file_end=out["datetime"].max()
    out.insert(0,"site",site); out.insert(1,"serial_hobo",info["serial"]); out.insert(2,"deployment",info["deployment"])
    out["source_file"]=str(path);out["file_end"]=file_end
    diffs=out["datetime"].sort_values().drop_duplicates().diff().dropna().dt.total_seconds()/60
    step=float(diffs.mode().iloc[0]) if len(diffs) else np.nan
    audit={"site":site,"file":str(path),"status":"ok_selected","serial":info["serial"],"deployment":info["deployment"],"start":out["datetime"].min(),"end":out["datetime"].max(),"n_rows":len(out),"native_resolution_minutes":step,"note":"largest cumulative file for deployment"}
    return out,audit


def process(root: Path, config_path: Path | None = None):
    cfg=load_config(config_path);min_light=int(cfg.get("light_period",{}).get("min_intervals_daily",20))
    selected,audits=select_cumulative_files(root)
    frames=[]
    print(f"Microclima: {len(selected)} arquivo(s) cumulativo(s) selecionado(s).", flush=True)
    for i,info in enumerate(selected, start=1):
        size_mb = info["path"].stat().st_size / (1024**2)
        print(
            f"  [{i}/{len(selected)}] {info['site']} | {info['path'].name} | {size_mb:.1f} MB",
            flush=True,
        )
        data,audit=read_one(info);audits.append(audit)
        if data is not None:
            frames.append(data)
            print(f"      {len(data):,} registros lidos.", flush=True)
    if not frames:raise RuntimeError("Nenhum RAW válido de Termohigrometro foi encontrado.")
    native=pd.concat(frames,ignore_index=True)
    native=native.sort_values(["site","datetime","file_end","source_file"],ascending=[True,True,False,False]).drop_duplicates(["site","datetime"],keep="first").sort_values(["site","datetime"])

    x=native.copy();x["datetime_30min"]=x["datetime"].dt.floor("30min")
    agg=x.groupby(["site","datetime_30min"],as_index=False).agg(
        temp_C=("temp_C","mean"),temp_min_C=("temp_C","min"),temp_max_C=("temp_C","max"),rh_pct=("rh_pct","mean"),
        vpd_kPa=("vpd_kPa","mean"),vpd_max_kPa=("vpd_kPa","max"),par_umol=("par_umol","mean"),dewpoint_C=("dewpoint_C","mean"),n_native=("datetime","count")
    ).rename(columns={"datetime_30min":"datetime"})
    agg["date"]=agg["datetime"].dt.date;agg["periodo_luz"]=(agg["datetime"].dt.hour>=6)&(agg["datetime"].dt.hour<18)

    luz=agg[agg["periodo_luz"]].copy();rows=[]
    for (site,date),g in luz.groupby(["site","date"]):
        def mean_if(col):
            z=g[col].dropna();return z.mean() if len(z)>=min_light else np.nan
        n_rh=int(g["rh_pct"].notna().sum());n_sat=int((g["rh_pct"].notna()&(g["rh_pct"]>=99.9)).sum());frac=n_sat/n_rh if n_rh else np.nan;n_par=int(g["par_umol"].notna().sum())
        rows.append({"date":date,"site":site,"n_light_temp":int(g["temp_C"].notna().sum()),"n_light_rh":n_rh,"n_light_vpd":int(g["vpd_kPa"].notna().sum()),"n_light_par":n_par,
                     "temp_luz_C":mean_if("temp_C"),"rh_luz_pct":mean_if("rh_pct"),"vpd_luz_kPa":mean_if("vpd_kPa"),"par_luz_umol":mean_if("par_umol"),
                     "dli_mol_m2":g["par_umol"].sum()*1800/1e6 if n_par>=min_light else np.nan,"n_rh_saturado":n_sat,"fracao_rh_saturado":frac,
                     "qa_rh_vpd":"sem_dados" if pd.isna(frac) else "ok_observed_condition", "rh_saturation_observed": bool(frac>0.50) if not pd.isna(frac) else False})
    daily=pd.DataFrame(rows)
    if not daily.empty:
        daily["rh_luz_qc_pct"]=daily["rh_luz_pct"];daily["vpd_luz_qc_kPa"]=daily["vpd_luz_kPa"]

    processed=root/"data"/"processed"/"microclima";results=root/"results"/"microclima";processed.mkdir(parents=True,exist_ok=True);results.mkdir(parents=True,exist_ok=True)
    native.drop(columns=["file_end"]).to_csv(processed/"microclima_native.csv",index=False,encoding="utf-8-sig")
    agg.to_csv(processed/"microclima_30min_long.csv",index=False,encoding="utf-8-sig");daily.to_csv(processed/"microclima_diario_06_18.csv",index=False,encoding="utf-8-sig")
    pd.DataFrame(audits).to_csv(results/"inventario_arquivos_microclima.csv",index=False,encoding="utf-8-sig")
    value_cols=["temp_C","temp_min_C","temp_max_C","rh_pct","vpd_kPa","vpd_max_kPa","par_umol","dewpoint_C"]
    wide=agg.pivot(index="datetime",columns="site",values=value_cols);wide.columns=[f"{site}_{var}" for var,site in wide.columns];wide=wide.reset_index().sort_values("datetime")
    wide.to_csv(processed/"microclima_30min_wide.csv",index=False,encoding="utf-8-sig")
    return {"native":native,"long30":agg,"wide30":wide,"daily":daily,"audits":pd.DataFrame(audits)}


def main():
    ap=argparse.ArgumentParser();ap.add_argument("--root",required=True);ap.add_argument("--config",default=None);a=ap.parse_args();r=process(Path(a.root),Path(a.config) if a.config else None)
    print(f"Microclima: {len(r['native']):,} registros nativos; {len(r['long30']):,} registros de 30 min.")
if __name__=="__main__":main()
